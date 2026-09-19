import copy, docx, os
HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "AushadhiNet_Review1_Document_v2.docx"
src = docx.Document(SRC)
survey = copy.deepcopy(src.tables[0]._tbl)
for t in survey.iter(qn("w:t")):
    if "29%" in (t.text or ""): print("fix:", t.text); t.text = t.text.replace("register completeness rose from 29% to 75%", "daily stock updating rose from 38.6% to 53.5% of facilities")
d = docx.Document(SRC)
body = d.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"): body.remove(el)

H1 = lambda t: d.add_heading(t, 1)
H2 = lambda t: d.add_heading(t, 2)
def P(t, bold=False, italic=False, align=None, size=None):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p
def B(items):
    for t in items:
        p = d.add_paragraph(style="List Paragraph")
        if isinstance(t, tuple):
            r = p.add_run("•  " + t[0] + " "); r.bold = True; p.add_run(t[1])
        else: p.add_run("•  " + t)
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexc); tcPr.append(s)
def T(header, rows, caption, widths=None):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid" if "Table Grid" in [s.name for s in d.styles] else None
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(9)
    borders(t)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    P(caption, italic=True, size=9.5)
def borders(t):
    tblPr = t._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top","left","bottom","right","insideH","insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"4"); x.set(qn("w:color"),"808080"); b.append(x)
    tblPr.append(b)
def FIG(path, cap):
    d.add_picture(path, width=Inches(6.5)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(cap, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)

# ---------- Title ----------
P("AushadhiNet", bold=True, size=26, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Federated AI for India's Last-Mile Health Resource Network — Medicines, Beds and Health Workforce", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Review – 1 Document (Revision 2)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Build with AI: Code for Communities — Second Edition", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Problem Statement 03 — Smart Health & Supply Chain Resilience", align=WD_ALIGN_PARAGRAPH.CENTER)
P("Components covered in this document:", bold=True)
B(["Abstract and one-line solution summary","Problem Statement, Motivation (figures verified against primary sources)",
   "Literature Survey, Survey Table, Gap Identification & Objectives","Design Methodology with Diagram",
   "Architecture — Modules, Description and Google Cloud technology mapping","Dataset Description",
   "Evaluation Protocol (held-out, leakage-controlled)","Deployability, Scalability across India and BRICS applicability",
   "Scope Statement and Known Limitations","References — APA 7th edition"])

H1("Abstract")
P("India's public health supply chain is digitised at the warehouse and thin at the point of care. The Drugs and Vaccine Distribution Management System (DVDMS, branded e-Aushadhi in most states) is deployed in more than twenty states and nominally reaches Primary Health Centres (PHCs), and the Electronic Vaccine Intelligence Network (eVIN) digitises vaccine cold-chain points. Yet audits and field studies report that data entry at peripheral facilities is incomplete and delayed, that sub-centres and many PHCs still rely on paper registers, and that neither bed availability nor health-worker attendance is visible in real time alongside stock. A stock-out is therefore discovered when a patient is turned away, while a neighbouring facility may hold surplus stock that expires unused, a vacant bed may exist one block away, or a posted doctor may be absent without the district knowing.")
P("AushadhiNet is a federated AI platform built on Google Cloud that closes this loop for the three resources named in the problem statement: medicines, beds and medical personnel. (1) A near-zero-friction capture layer turns a WhatsApp voice note, a photograph of a handwritten register, a geo-tagged staff check-in or an IVR call into structured, confidence-scored records using Chirp speech recognition and Gemini multimodal extraction, with a mandatory human confirmation for any low-confidence field. (2) A hierarchical forecasting stack (TimesFM, gradient-boosted models on Vertex AI and intermittent-demand estimators) predicts facility × drug stock-out probability, bed occupancy and staffing gaps two to six weeks ahead. (3) An OR-Tools optimisation module converts forecasts into signed, expiry-aware, drive-time-feasible transfer orders, staff deputation suggestions and patient-referral recommendations, which a district officer approves through a Gemini function-calling agent in their own language. (4) A federation layer lets states improve a shared forecasting model by exchanging differentially private model updates instead of raw facility data.")
P("Evaluation is designed before any model is trained: forecasting accuracy is reported on a frozen, time-based held-out window and on an entire held-out district and state; results on real public data (HMIS service-delivery series) and on a sealed synthetic stock ledger are reported separately and never combined; and the synthetic generator uses hidden drivers that the forecasting model cannot observe, so that the model cannot score well simply by inverting the generator.")
P("One-line summary: ", bold=True).add_run("AushadhiNet lets a nurse report medicine stock, beds and staff presence by WhatsApp voice or photo in any Indian language, predicts shortages weeks ahead with Google AI, and turns each prediction into an approved transfer order that moves medicine, not just a dashboard alert.").bold = False

# ---------- 1 ----------
H1("1. Problem Statement")
P("DVDMS/e-Aushadhi, developed by C-DAC, automates procurement, warehousing and issue of drugs and is reported in use across more than twenty states, covering district warehouses, hospitals, Community Health Centres (CHCs) and, in several states, PHCs. The Government of India has asked states to extend DVDMS down to the drug-distribution-counter level, which is itself evidence that last-mile coverage is not yet complete. Performance audits of e-Aushadhi (for example, the Comptroller and Auditor General's audit for Punjab) report a consistent set of unresolved deficiencies: shortage of dedicated data-entry staff at peripheral facilities, incomplete adoption of barcoding, lack of dependable connectivity, and continuing stock-outs and expiry despite the digitised upstream pipeline.")
P("The consequence is a visibility gap that is widest exactly where the citizen receives care. At the 1,69,615 Sub-Centres and in many of the 31,882 PHCs recorded in Health Dynamics of India 2022-23, stock is maintained or first recorded on paper, transcribed late or not at all, and consumed data arrives in monthly aggregates. Bed occupancy and staff attendance at these facilities are not linked to the stock picture at all, even though the problem statement names all three.")
P("The problem is therefore three-fold:")
B([("Capture:","there is no low-friction way for a non-data-entry health worker to report stock, beds and staff presence at the frequency and quality needed for prediction, in their own language, on a basic phone."),
   ("Prediction:","even where data exists, systems report what has happened; none forecasts facility-level stock-out, bed pressure or staffing gaps early enough to act."),
   ("Action and governance:","no system converts visibility into an executable, constraint-aware redistribution action between facilities, and any national solution must respect that each state owns its health data.")])
P("Problem addressed: how can India obtain reliable, low-latency, last-mile visibility of medicines, beds and personnel across a linguistically diverse, connectivity-constrained PHC network, convert that visibility into approved cross-facility redistribution decisions, and allow states to share predictive models without surrendering facility-level data?", italic=True)

H1("2. Motivation")
P("Scale. Health Dynamics of India (Infrastructure & Human Resources) 2022-23 records 1,69,615 Sub-Centres, 31,882 PHCs and 6,359 CHCs as of 31 March 2023. The overwhelming majority of front-line reporting is done by Auxiliary Nurse Midwives and pharmacists who are not data-entry professionals and who work under variable connectivity.")
P("Evidence that reducing friction works. India's own eVIN programme shows what happens when a facility touchpoint is digitised: the programmatic assessment by Gurnani et al. (2020) across 12 states found that the number of cold-chain facilities with a stock-out of any vaccine fell by 30.4% after eVIN (p < 0.001), the share of facilities updating vaccine stock daily rose from 38.6% to 53.5%, and completeness of stock registers, indent forms and temperature logbooks improved significantly. AushadhiNet applies the same principle, remove friction and compliant data follows, to the general drug register, beds and attendance, and adds the forecasting and redistribution steps eVIN does not perform.")
P("Technology readiness. The components now exist at production quality on a single cloud: Google's Chirp/USM speech models cover 100+ languages including major Indian languages; Gemini accepts images and returns schema-constrained JSON; TimesFM provides a pretrained time-series foundation model useful for short-history facilities; Vertex AI and BigQuery ML serve custom models at national scale; and Google OR-Tools solves min-cost-flow and routing problems in milliseconds for district-sized networks. Federated learning frameworks such as Flower have matured to the point where cross-institution training can run on managed infrastructure.")
P("Policy fit. India's Digital Public Infrastructure approach (Aadhaar, UPI, ABDM) succeeds when systems are open, interoperable and privacy-preserving by design. Health is a State subject; a federated design in which raw facility data stays inside each state's own cloud project is the architecture most consistent with how Indian health data is actually governed, and therefore the most adoptable.")

# ---------- 3 ----------
H1("3. Literature Survey and Gap Identification & Objectives")
H2("3.1 Literature Survey")
P("The relevant literature spans five largely separate communities: (i) machine-learning demand forecasting in pharmaceutical and health settings; (ii) federated and privacy-preserving learning in healthcare; (iii) handwritten and Indic-script document recognition; (iv) conversational and voice AI for low-resource, multilingual health settings; and (v) vehicle-routing and network-flow optimisation for health logistics. India-specific programmatic literature documents the operational reality of e-Aushadhi/DVDMS, eVIN and HMIS. Sixteen representative studies are compared in Table 1.")
P("Forecasting studies (rows 1, 14) show that gradient-boosted and tree-based models beat naive baselines when given rich covariates, but are evaluated on manufacturer sales or surveillance data, not facility-level public-health stock, and rarely address the intermittent, zero-heavy demand typical of a PHC. Federated-learning studies (rows 3, 4) establish FedAvg and report that healthcare FL is simulated far more often than deployed. OCR studies (rows 5–7) reach high accuracy on isolated Devanagari characters or constrained word sets but stop short of structured, confidence-scored multi-field extraction from degraded, mixed-script registers. Conversational-AI studies (rows 8, 9, 13, 16) answer questions but do not drive an inventory decision. Optimisation studies (rows 10–12) are mature but framed around disaster relief rather than routine, expiry-ordered lateral transfer. India-specific studies (rows 2, 15) confirm both the benefit of digitising a touchpoint and the persistence of connectivity and workload constraints.")
H2("3.2 Survey Table")
_ph = d.add_paragraph(); _ph._p.addnext(survey)
P("Table 1. Structured comparison of sixteen representative studies across the five research threads underlying AushadhiNet.", italic=True, size=9.5)
H2("3.3 Research Gaps Identified")
B([("Gap 1 — Fragmentation.","Capture, forecasting, optimisation and conversational AI are studied separately; no reviewed system connects last-mile capture to forecasting to an executable transfer action."),
   ("Gap 2 — Structured, confidence-aware extraction.","No reviewed OCR system performs multi-field, confidence-scored extraction (drug, batch, quantity, expiry) from real, degraded, mixed-script registers with a confirmation fallback."),
   ("Gap 3 — Federated learning remains simulated.","Healthcare FL is rarely deployed across institutions, and no reviewed study applies it to cross-state pharmaceutical demand forecasting."),
   ("Gap 4 — Indian systems stop at the wrong tier and the wrong resource.","DVDMS and eVIN digitise warehouse and cold-chain transactions and HMIS aggregates monthly indicators, but none gives near-real-time stock, bed and attendance visibility at PHC/sub-centre level or couples it to prediction."),
   ("Gap 5 — Optimisation framed around disasters.","Routing literature targets post-disaster relief, not recurring, expiry-ordered, drive-time-constrained lateral transfers between peer PHCs, nor joint drug-and-staff rebalancing."),
   ("Gap 6 — Conversational AI is not decision-facing.","No reviewed system lets an officer convert a natural-language request into an approvable, auditable redistribution order."),
   ("Gap 7 — Evaluation practice.","Forecasting studies on synthetic or proprietary data rarely separate the data-generating process from the model or report on held-out geography, making reported gains hard to trust.")])
H2("3.4 Objectives")
B(["O1. Build a multilingual voice-, photo-, IVR- and check-in-based capture pipeline that converts handwritten registers, bed census and staff presence into structured, confidence-scored records (Gaps 2, 4).",
   "O2. Build a hierarchical facility × drug × week forecasting model fusing consumption with HMIS service-delivery, IDSP/IHIP outbreak and IMD weather signals, producing calibrated stock-out probability at 2, 4 and 6 weeks, plus bed-occupancy and staffing-gap forecasts (Gap 4).",
   "O3. Formulate and solve a constrained redistribution problem that outputs expiry-ordered, drive-time-feasible drug transfers, temporary staff deputation and patient-referral suggestions (Gap 5).",
   "O4. Implement a federated-learning protocol across state data planes with FedProx aggregation and differential privacy, and measure the accuracy gain for a data-poor state against local-only and centralised training (Gap 3).",
   "O5. Build a Gemini function-calling agent through which a block or district officer can query system state and approve actions in natural language, with every action logged (Gaps 1, 6).",
   "O6. Evaluate the integrated system under a pre-registered protocol with frozen held-out time windows and held-out geography, separate reporting of real and synthetic results, and fixed baselines (Gap 7).",
   "O7. Publish an open ingestion contract and adapters (DVDMS export, HMIS, ABDM Health Facility Registry, FHIR) so that a new state or country can onboard without migrating its existing system."])

# ---------- 4 ----------
H1("4. Design Methodology")
P("The project follows a nine-stage methodology. The key change from the first revision is Stage 4: the evaluation protocol, data splits and synthetic-data generator are fixed and committed before any model is designed or tuned, so that no reported number can be the product of tuning against the test set or of a model learning the generator's formula.")
FIG(HERE + "/fig1_methodology.png", "Figure 1. Nine-stage design methodology for AushadhiNet. Highlighted stages protect the integrity of reported results.")
B([("Stage 1 — Problem grounding:","hackathon brief, CAG audit, Health Dynamics of India 2022-23, DVDMS and eVIN literature."),
   ("Stage 2 — Literature survey:","Section 3."),
   ("Stage 3 — Requirements:","functional (capture, forecast, redistribute, federate, converse, for medicines, beds and staff) and non-functional (offline tolerance, low per-facility cost, state data sovereignty, interpretability for non-technical officers, language coverage)."),
   ("Stage 4 — Evaluation protocol frozen first:","splits, metrics, baselines and the sealed synthetic generator are committed with a content hash before modelling (Section 7)."),
   ("Stage 5 — Data acquisition:","real public datasets and the synthetic ledger (Section 6)."),
   ("Stages 6–8 — Design, implementation, integration:","the five modules of Section 5, integrated into one end-to-end loop."),
   ("Stage 9 — Held-out evaluation:","the test window and held-out geography are scored once, and results are reported under the rules of Section 7.")])

# ---------- 5 ----------
H1("5. Architecture — Modules and Description")
P("AushadhiNet has five layered modules on Google Cloud. Data flows downward from capture to action; alerts and approvals flow back upward through the interface layer.")
FIG(HERE + "/fig2_architecture.png", "Figure 2. AushadhiNet layered architecture and the Google Cloud services used in each layer.")
H2("5.1 Module 1 — Interface and Capture")
P("The only point of contact for facility staff, designed around the phone they already carry. Channels: (a) WhatsApp Business API — voice note or photograph of the stock register, a daily one-line bed census, and a geo-tagged staff check-in validated against the facility's ABDM Health Facility Registry coordinates; (b) an IVR line built on Dialogflow CX telephony for feature phones and no-data areas; (c) an offline-first Progressive Web App on Firebase Hosting for block and district officers; and (d) a map-first dashboard on Google Maps Platform with drill-down from national to facility level. Where a state already runs a biometric attendance system or DVDMS, the module ingests that feed instead of duplicating it.")
H2("5.2 Module 2 — Ingestion and Extraction")
P("Voice notes are transcribed by Chirp 2 through Cloud Speech-to-Text v2 and passed to Gemini with a response schema to extract drug, quantity, batch and expiry (or bed and staff fields). Register photographs go directly to Gemini multimodal, which returns schema-constrained JSON with a per-field confidence. Drug names are normalised against the National List of Essential Medicines using embeddings in Vertex AI Vector Search, resolving brand names, generics and local-language spellings. No field below the confidence threshold is written silently: it triggers a one-tap confirmation card to the reporter. Records are published to Pub/Sub, written to Firestore for live state and streamed to BigQuery for history. An open, versioned ingestion contract allows a DVDMS export or state HMIS feed to be mapped into the same schema.")
H2("5.3 Module 3 — Intelligence")
B([("Demand forecasting:","a hierarchical facility × drug × week model. TimesFM provides zero-shot forecasts for facilities with short histories; a LightGBM model trained on Vertex AI uses HMIS footfall, IDSP/IHIP outbreak signals, IMD weather and programme/festival calendars; Croston/SBA handles intermittent, zero-heavy drugs; BigQuery ML ARIMA_PLUS serves as a strong statistical baseline. A stacked ensemble outputs quantile forecasts that are converted into calibrated stock-out probabilities at 2, 4 and 6 weeks, and forecasts are reconciled across facility, block and district levels."),
   ("Bed and staffing forecasts:","bed-occupancy forecasts from the daily census and HMIS inpatient series; staffing-gap forecasts from attendance history, sanctioned posts and expected footfall."),
   ("Anomaly detection:","flags consumption or attendance patterns not explained by known drivers (possible data error, leakage or an undetected outbreak) for human review."),
   ("Redistribution solver:","Google OR-Tools min-cost-flow across a cluster of facilities. Objective: minimise expected stock-out days plus expiry wastage plus transport cost. Constraints: real drive times from Google Maps, cold-chain eligibility, minimum retention buffers and first-expiry-first-out batch ordering. The same solver proposes temporary staff deputation and patient referral to facilities with free beds. The output is a concrete order, not an alert."),
   ("Officer agent:","a Gemini function-calling agent built with Google's Agent Development Kit. Tools are read-only queries on BigQuery and Firestore plus a solver call; the agent drafts orders but can never execute one without explicit officer approval, and every tool call is audit-logged.")])
H2("5.4 Module 4 — Federation")
P("Each participating state runs its own data plane in its own Google Cloud project; raw facility rows never leave it. A central model plane runs a Flower federated server on Vertex AI. States train locally and send clipped, noised (DP-SGD) model updates, aggregated with FedProx to cope with non-identical data across states and protected by secure aggregation. The improved global model is returned to every state, so a state with a short data history benefits from others without transferring facility-identifiable data. The layer also maintains the published ingestion contract, so onboarding a new state means writing an adapter, not migrating a system.")
H2("5.5 Module 5 — Action and Delivery")
P("Turns an approved decision into executable artefacts: a transfer order signed with Cloud KMS and carrying a batch-level manifest; WhatsApp and Text-to-Speech voice alerts to the source and destination in-charges in their language; route plans from the Google Maps Routes API; a time-boxed escalation ladder (facility → block → district → state) orchestrated by Cloud Workflows and Cloud Scheduler; and a public, aggregate-only transparency view of district stock-out days in Looker Studio.")
H2("5.6 Google Cloud and Google AI Technology Mapping")
T(["Capability","Google technology","Role in AushadhiNet"],[
 ["Speech recognition","Chirp 2 via Cloud Speech-to-Text v2","Voice notes and IVR in Indian languages"],
 ["Multimodal extraction","Gemini (Vertex AI), structured output","Handwritten register photo to JSON with per-field confidence"],
 ["Translation / voice out","Cloud Translation, Cloud Text-to-Speech","Local-language alerts and confirmation cards"],
 ["IVR / dialogue","Dialogflow CX (telephony)","Feature-phone reporting"],
 ["Entity normalisation","Vertex AI embeddings + Vector Search","Map extracted drug names to NLEM"],
 ["Forecasting","TimesFM, Vertex AI custom training, BigQuery ML ARIMA_PLUS","Stock-out, bed and staffing forecasts"],
 ["Optimisation","Google OR-Tools","Min-cost-flow transfers, deputation, referral"],
 ["Agent","Gemini function calling, Agent Development Kit","Officer queries and order drafting"],
 ["Federated learning","Flower on Vertex AI, one GCP project per state","Cross-state model sharing without raw data"],
 ["Data","Firestore, BigQuery, Pub/Sub","Live state, history, streaming"],
 ["Geospatial","Maps JavaScript API, Distance Matrix / Routes API","Dashboard, drive-time constraints, routes"],
 ["Serving & security","Cloud Run, Firebase, IAM, VPC Service Controls, Cloud KMS, Audit Logs","Hosting, isolation, signing, audit"],
 ["Reporting","Looker Studio","Public aggregate transparency"]],
 "Table 2. Mapping of every AushadhiNet capability to a Google Cloud / Google AI service.", widths=[1.5,2.4,2.6])

# ---------- 6 ----------
H1("6. Dataset Description")
H2("6.1 Real, publicly sourced datasets")
B([("HMIS (MoHFW):","facility-level monthly service-delivery indicators (OPD footfall, deliveries, immunisation sessions, ANC visits, inpatient admissions) — demand proxy and the source of the real-data forecasting benchmark."),
   ("ABDM Health Facility Registry:","facility identifiers, type and geocoded location; every record is anchored to a verifiable facility."),
   ("Health Dynamics of India (Infrastructure & HR) 2022-23:","facility counts, sanctioned and in-position staff, used to size pilot districts and as the baseline for staffing-gap estimates."),
   ("IDSP/IHIP:","weekly outbreak bulletins, used as an epidemiological feature and a surge re-forecast trigger."),
   ("IMD:","rainfall, temperature and humidity for vector-borne and diarrhoeal seasonality (ORS, zinc, antimalarials)."),
   ("NLEM 2022:","drug master list for name normalisation and cold-chain flags."),
   ("Census district demographics:","catchment population weights."),
   ("Google Maps Platform:","real drive times between facilities as hard solver constraints.")])
H2("6.2 Sealed synthetic stock, bed and attendance ledger")
P("Per-facility, per-drug daily stock ledgers, daily bed census and attendance logs are not publicly released, so a seeded generator produces them. To prevent the forecasting model from simply learning the generator back (an 'oracle' advantage), the generator is built under four rules:")
B(["Hidden drivers: demand is driven partly by latent factors the model never sees — unannounced outbreak shocks, supply lead-time disruptions, leakage and reporting delays — in addition to the observable HMIS/IDSP/IMD covariates.",
   "Sealed parameters: generator code and parameters live in a separate module whose content hash and random seed are committed before any forecasting code is written; the forecasting code may not import it.",
   "Realistic noise: late entries, missing weeks, transcription errors and intermittent zero-demand are injected.",
   "Full publication: generator, seed and transformation logic are released so every synthetic result is reproducible, and every synthetic figure is labelled as synthetic."])
H2("6.3 Handwriting-extraction evaluation set")
P("A separately labelled set of at least 200 photographs of stock-register-style pages in Devanagari, Tamil and Bengali script, including degraded samples (strikethrough, faded ink, mixed script, skew, low light). Each page is labelled independently by two annotators and inter-annotator agreement (Cohen's kappa) is reported. The set is split into a prompt-development portion and a held-out test portion before any prompt is tuned.")

# ---------- 7 ----------
H1("7. Evaluation Protocol")
P("The protocol is committed before modelling (Stage 4) and follows established out-of-sample practice (Tashman, 2000) and data-leakage guidance (Kaufman et al., 2012).")
T(["Component","Held-out design","Metrics","Baselines"],[
 ["Forecasting (real HMIS)","Rolling-origin; final 26 weeks frozen as test; one entire district never seen in training","WAPE, MASE, pinball loss (quantiles)","Seasonal naive, moving average, ARIMA_PLUS"],
 ["Stock-out prediction (synthetic)","Same time split; one held-out district and one held-out state","Recall and precision at 4 weeks, Brier score, calibration curve","Moving average + reorder point"],
 ["Redistribution","Replay on test window only","Stock-out days, expired units, km travelled","Status-quo monthly indent policy"],
 ["Federated learning","States as clients; data-poor state evaluated on its own held-out weeks","Accuracy gain vs local-only; gap vs centralised","Local-only, centralised upper bound"],
 ["Extraction","Held-out 50%+ of labelled register set","Field-level exact match, character error rate, confidence calibration","Cloud Vision OCR + rules"],
 ["End-to-end","Live demo facility loop","p50 / p95 capture-to-record latency","—"]],
 "Table 3. Pre-registered evaluation protocol.", widths=[1.4,2.0,1.7,1.4])
P("Reporting rules: (1) MAPE is not used as the headline metric because PHC demand contains many zero weeks, where MAPE is undefined; WAPE and MASE are used instead. (2) Real-data and synthetic-data results are reported in separate tables and never averaged together. (3) The test window and held-out geography are scored once; any later change to the model is reported as a new, labelled run. (4) The headline claim of the project is taken from the real-data benchmark, and synthetic results are presented as supporting evidence for the redistribution logic only.")

# ---------- 8 ----------
H1("8. Deployability, Scalability and Reach")
B([("Pilot in weeks:","no new hardware; WhatsApp and IVR reach any phone; fully serverless (Cloud Run, Firestore, BigQuery) so a new district is configuration, not infrastructure."),
   ("State onboarding:","one Terraform module creates a state's data plane; an adapter maps its DVDMS or HMIS export to the open ingestion contract."),
   ("National scale:","BigQuery and Vertex AI scale to all ~2 lakh reporting facilities; the solver runs per district cluster, so cost grows linearly."),
   ("Language reach:","Chirp, Gemini and Cloud Translation cover the major scheduled languages; adding a language is configuration."),
   ("Interoperability:","ABDM Health Facility Registry identifiers and an optional FHIR R4 export; designed as a Digital Public Good with open schema and open-source adapters."),
   ("BRICS and cross-border applicability:","the ingestion contract is keyed to a configurable essential-medicines list and facility registry, so the same stack can run in another country by swapping the national list, registry and languages; the federated design lets countries share models without sharing data.")])

# ---------- 9 ----------
H1("9. Scope Statement and Known Limitations")
B(["Federated learning is demonstrated with states as separate clients in separate cloud projects using real, state-partitioned HMIS data plus the synthetic ledger. It is a working protocol, not a live multi-government deployment, which would require data-sharing agreements outside the scope of a hackathon.",
   "Stock-out and redistribution results rely on the synthetic ledger because facility-level stock ledgers are not public; they validate logic, not real-world impact. A field pilot with a state's DVDMS data is the next step.",
   "Staff attendance uses geo-tagged self check-in, which can be gamed; it is positioned as a visibility aid that complements, not replaces, official biometric systems.",
   "The system recommends; a human officer always approves any transfer, deputation or referral."])

# ---------- References ----------
H1("References")
refs = [
"Zhu, X., Ninh, A., Zhao, H., & Liu, Z. (2021). Demand forecasting with supply-chain information and machine learning: Evidence in the pharmaceutical industry. Production and Operations Management, 30(9), 3231–3252. https://doi.org/10.1111/poms.13426",
"Gurnani, V., Singh, P., Haldar, P., Aggarwal, M. K., Agrahari, K., Kashyap, S., Ghosh, S., Mohapatra, M. K., Bhargava, R., Nandi, P., Dhalaria, P., & Wai, K. T. (2020). Programmatic assessment of electronic Vaccine Intelligence Network (eVIN). PLOS ONE, 15(11), Article e0241369. https://doi.org/10.1371/journal.pone.0241369",
"Zerka, F., Barakat, S., Walsh, S., Bogowicz, M., Leijenaar, R. T. H., Jochems, A., Miraglio, B., Townend, D., & Lambin, P. (2020). Systematic review of privacy-preserving distributed machine learning from federated databases in health care. JCO Clinical Cancer Informatics, 4, 184–200. https://doi.org/10.1200/CCI.19.00047",
"McMahan, H. B., Moore, E., Ramage, D., Hampson, S., & y Arcas, B. A. (2017). Communication-efficient learning of deep networks from decentralized data. Proceedings of AISTATS, 54, 1273–1282.",
"Dey, S., Dutta, M. K., & Nasipuri, M. (2021). An efficient CNN-based deep learning model for handwritten Devanagari character recognition. International Journal of Machine Learning and Cybernetics, 12, 89–103. https://doi.org/10.1007/s13042-020-01152-3",
"Roy, R. K., Mukherjee, H., Roy, K., & Pal, U. (2022). CNN based recognition of handwritten multilingual city names. Multimedia Tools and Applications, 81(8), 11501–11517. https://doi.org/10.1007/s11042-021-11914-5",
"Guha, R., Das, N., Kundu, M., Nasipuri, M., & Santosh, K. C. (2020). DevNet: An efficient CNN architecture for handwritten Devanagari character recognition. International Journal of Pattern Recognition and Artificial Intelligence, 34(12), Article 2052009.",
"Mishra, R., Singh, S., Kaur, J., Singh, P., & Shah, R. (2023). Hindi chatbot for supporting maternal and child health related queries in rural India. In Proceedings of the 5th Clinical Natural Language Processing Workshop (pp. 69–77). ACL.",
"Radford, A., Kim, J. W., Xu, T., Brockman, G., McLeavey, C., & Sutskever, I. (2023). Robust speech recognition via large-scale weak supervision. Proceedings of ICML, 202, 28492–28518.",
"Maroof, A., Khalid, Q. S., Mahmood, M., Naeem, K., Maqsood, S., Khattak, S. B., & Ayvaz, B. (2023). Vehicle routing optimization for humanitarian supply chain: A systematic review of approaches and solutions. IEEE Access, 11, 127157–127175. https://doi.org/10.1109/ACCESS.2023.3331062",
"Nodoust, S., Pishvaee, M. S., & Seyedhosseini, S. M. (2023). Vehicle routing problem for humanitarian relief distribution under hybrid uncertainty. Kybernetes, 52(4), 1503–1527. https://doi.org/10.1108/K-09-2021-0839",
"Ahuja, R. K., Magnanti, T. L., & Orlin, J. B. (1993). Network flows: Theory, algorithms, and applications. Prentice Hall.",
"Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., Küttler, H., Lewis, M., Yih, W., Rocktäschel, T., Riedel, S., & Kiela, D. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems, 33, 9459–9474.",
"Rahman, M. S., Amrin, M., & Shiddik, M. A. B. (2025). Dengue early warning system and outbreak prediction tool in Bangladesh using interpretable tree-based machine learning model. Health Science Reports, 8(5), Article e70726. https://doi.org/10.1002/hsr2.70726",
"Chaudhary, K., Awasthi, S., Kanubhai, T. H., & Preeti. (2025). Assessment of eVIN (Electronic Vaccine Intelligence Network) in utility and constraints of various stakeholders in Nainital District: A mixed method study. Journal of Family Medicine and Primary Care, 14(2), 627–633.",
"Huang, Z., Xue, K., Fan, Y., Mu, L., Liu, R., Ruan, T., Zhang, S., & Zhang, X. (2024). Tool calling: Enhancing medication consultation via retrieval-augmented large language models (arXiv:2404.17897). arXiv.",
"Gemini Team, Google. (2023). Gemini: A family of highly capable multimodal models (arXiv:2312.11805). arXiv. https://doi.org/10.48550/arXiv.2312.11805",
"Zhang, Y., Han, W., Qin, J., Wang, Y., Bapna, A., Chen, Z., et al. (2023). Google USM: Scaling automatic speech recognition beyond 100 languages (arXiv:2303.01037). arXiv.",
"Das, A., Kong, W., Sen, R., & Zhou, Y. (2024). A decoder-only foundation model for time-series forecasting. Proceedings of the 41st International Conference on Machine Learning (ICML), PMLR 235.",
"Perron, L., & Furnon, V. (2024). OR-Tools (Version 9) [Computer software]. Google. https://developers.google.com/optimization/",
"Beutel, D. J., Topal, T., Mathur, A., Qiu, X., Fernandez-Marques, J., Gao, Y., et al. (2020). Flower: A friendly federated learning framework (arXiv:2007.14390). arXiv.",
"Li, T., Sahu, A. K., Zaheer, M., Sanjabi, M., Talwalkar, A., & Smith, V. (2020). Federated optimization in heterogeneous networks. Proceedings of Machine Learning and Systems, 2, 429–450.",
"Kairouz, P., McMahan, H. B., Avent, B., et al. (2021). Advances and open problems in federated learning. Foundations and Trends in Machine Learning, 14(1–2), 1–210.",
"Abadi, M., Chu, A., Goodfellow, I., McMahan, H. B., Mironov, I., Talwar, K., & Zhang, L. (2016). Deep learning with differential privacy. Proceedings of the 2016 ACM SIGSAC Conference on Computer and Communications Security, 308–318. https://doi.org/10.1145/2976749.2978318",
"Dwork, C., & Roth, A. (2014). The algorithmic foundations of differential privacy. Foundations and Trends in Theoretical Computer Science, 9(3–4), 211–407. https://doi.org/10.1561/0400000042",
"Croston, J. D. (1972). Forecasting and stock control for intermittent demands. Operational Research Quarterly, 23(3), 289–303.",
"Syntetos, A. A., & Boylan, J. E. (2005). The accuracy of intermittent demand estimates. International Journal of Forecasting, 21(2), 303–314.",
"Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q., & Liu, T.-Y. (2017). LightGBM: A highly efficient gradient boosting decision tree. Advances in Neural Information Processing Systems, 30, 3146–3154.",
"Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD, 785–794. https://doi.org/10.1145/2939672.2939785",
"Lim, B., Arık, S. Ö., Loeff, N., & Pfister, T. (2021). Temporal fusion transformers for interpretable multi-horizon time series forecasting. International Journal of Forecasting, 37(4), 1748–1764. https://doi.org/10.1016/j.ijforecast.2021.03.012",
"Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
"Tashman, L. J. (2000). Out-of-sample tests of forecasting accuracy: An analysis and review. International Journal of Forecasting, 16(4), 437–450.",
"Kaufman, S., Rosset, S., Perlich, C., & Stitelman, O. (2012). Leakage in data mining: Formulation, detection, and avoidance. ACM Transactions on Knowledge Discovery from Data, 6(4), Article 15. https://doi.org/10.1145/2382577.2382579",
"Ministry of Health & Family Welfare, Government of India. (2024). Health dynamics of India (Infrastructure & human resources) 2022-23. MoHFW.",
"Ministry of Health & Family Welfare, Government of India. (2018). Techno-economic assessment of Electronic Vaccine Intelligence Network (eVIN). MoHFW.",
"Comptroller and Auditor General of India. (2019). Report No. 4 of 2019 — Social, General and Economic Sectors (Non-PSUs), Government of Punjab, Chapter II: Performance audit of e-Aushadhi. Office of the CAG of India.",
"Centre for Development of Advanced Computing. (n.d.). Drug and Vaccine Distribution Management System (DVDMS) [Project documentation]. C-DAC. https://www.cdac.in/",
"National Health Systems Resource Centre. (n.d.). Letter to states for implementation of DVDMS up to drug distribution counter level. NHSRC. https://qps.nhsrcindia.org/node/22297",
"National Health Authority. (n.d.). Ayushman Bharat Digital Mission — Health Facility Registry. https://facility.abdm.gov.in/",
"Ministry of Health & Family Welfare, Government of India. (2022). National List of Essential Medicines 2022. MoHFW.",
]
for i, r in enumerate(refs, 1):
    p = d.add_paragraph(f"[{i}]  {r}"); p.paragraph_format.space_after = Pt(4)
    for run in p.runs: run.font.size = Pt(9.5)
P("Entries 5–7, 15 and 37–40 should be re-checked against the publisher record for exact page numbers or URLs before final submission.", italic=True, size=9)

d.save("AushadhiNet_Review1_Document_v2.docx")
print("saved", len(refs), "refs")
